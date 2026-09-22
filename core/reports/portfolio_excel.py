"""
PORTFOLIO EXCEL REPORT GENERATOR MODULE
=======================================
Generates multi-sheet styled Excel (.xlsx) workbooks for portfolio treasury audits.
Includes Executive Summary, Active Holdings Valuation, Trade Journal & Attribution,
and Historical Equity Snapshots with institutional typography, palette, and formatting.
"""

import io
import datetime
from typing import Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from database import Portfolio, Position, Trade, PortfolioSnapshot, SignalExecutionAttribution, HorusExecution
from core import TimeUtils


def build_portfolio_excel_report(portfolio_id: int) -> bytes:
    """Generate a 4-sheet styled Excel workbook from a portfolio database record.

    Sheets:
      1. Executive Summary: Treasury balances, risk metrics, MTM vs. Realized PnL.
      2. Holdings & Valuation: Active positions, cost basis, live value, risk status, targets.
      3. Trade Journal: Settled trades history, execution lanes, realized returns.
      4. Equity Snapshots: Historical time-series records of equity & cash.

    Returns:
        bytes: Raw XLSX binary data ready for HTTP response or file storage.
    """
    portfolio = Portfolio.get_or_none(Portfolio.id == portfolio_id)
    if not portfolio:
        raise ValueError(f"Portfolio #{portfolio_id} not found")

    wb = openpyxl.Workbook()
    if wb.active is not None:
        wb.remove(wb.active)

    # -------------------------------------------------------------------------
    # Styling Definitions
    # -------------------------------------------------------------------------
    font_family = "Segoe UI"
    title_font = Font(name=font_family, size=16, bold=True, color="0F172A")
    subtitle_font = Font(name=font_family, size=10, italic=True, color="64748B")
    section_font = Font(name=font_family, size=12, bold=True, color="1E293B")
    bold_font = Font(name=font_family, size=10, bold=True, color="0F172A")
    regular_font = Font(name=font_family, size=10, color="334155")
    mono_font = Font(name="Consolas", size=9, color="334155")

    header_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    sub_header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")

    green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    green_font = Font(name=font_family, size=10, bold=True, color="15803D")

    red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    red_font = Font(name=font_family, size=10, bold=True, color="B91C1C")

    amber_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    amber_font = Font(name=font_family, size=10, bold=True, color="B45309")

    blue_fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
    blue_font = Font(name=font_family, size=10, bold=True, color="0369A1")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    thin_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    def apply_headers(ws, row_idx: int, col_count: int, fill=header_fill):
        for col in range(1, col_count + 1):
            cell = ws.cell(row=row_idx, column=col)
            cell.font = header_font
            cell.fill = fill
            cell.alignment = align_center
            cell.border = thin_border

    def auto_fit_columns(ws, max_cols: Optional[int] = None):
        ws.views.sheetView[0].showGridLines = True
        cols = range(1, (max_cols or ws.max_column) + 1)
        for col in cols:
            max_len = 0
            col_letter = get_column_letter(col)
            for row in range(1, ws.max_row + 1):
                val = ws.cell(row=row, column=col).value
                if val is not None:
                    max_len = max(max_len, len(str(val)))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # -------------------------------------------------------------------------
    # Fetch Data
    # -------------------------------------------------------------------------
    positions = list(Position.select().where(Position.portfolio == portfolio.id).order_by(Position.ticker.asc(), Position.id.asc()))
    open_positions = [p for p in positions if str(p.status or "OPEN").upper() == "OPEN"]
    trades = list(Trade.select().where(Trade.portfolio == portfolio.id).order_by(Trade.exit_date.desc(), Trade.id.desc()))
    snapshots = list(PortfolioSnapshot.select().where(PortfolioSnapshot.portfolio == portfolio.id).order_by(PortfolioSnapshot.date.desc()))

    # Attributions mapping for trades
    trade_ids = [t.id for t in trades]
    attributions = list(SignalExecutionAttribution.select().where(SignalExecutionAttribution.execution.in_(
        [e.id for e in HorusExecution.select().where(HorusExecution.trade_id.in_(trade_ids))]
    ))) if trade_ids else []
    
    execution_by_trade_id = {}
    if trade_ids:
        for exec_row in HorusExecution.select().where(HorusExecution.trade_id.in_(trade_ids)):
            execution_by_trade_id[exec_row.trade_id] = exec_row

    # Financial Calculations
    cash_egp = float(portfolio.cash_egp or 0.0)
    cash_usd = float(portfolio.cash_usd or 0.0)

    total_cost_basis = 0.0
    total_market_value = 0.0
    floating_mtm_pnl = 0.0

    for pos in open_positions:
        entry = float(pos.entry_price or 0.0)
        shares = int(pos.shares or 0)
        current = float(pos.current_price if pos.current_price is not None else entry)
        cost = entry * shares
        val = current * shares
        total_cost_basis += cost
        total_market_value += val
        floating_mtm_pnl += (val - cost)

    net_worth_egp = cash_egp + total_market_value
    cash_drag_pct = (cash_egp / net_worth_egp * 100.0) if net_worth_egp > 0 else 0.0
    invested_pct = (total_market_value / net_worth_egp * 100.0) if net_worth_egp > 0 else 0.0

    settled_trades_count = len(trades)
    winning_trades = [t for t in trades if float(t.pnl or 0.0) > 0]
    losing_trades = [t for t in trades if float(t.pnl or 0.0) <= 0]
    win_rate = (len(winning_trades) / settled_trades_count * 100.0) if settled_trades_count > 0 else 0.0
    settled_realized_pnl = sum(float(t.pnl or 0.0) for t in trades)

    gross_profit = sum(float(t.pnl or 0.0) for t in winning_trades)
    gross_loss = abs(sum(float(t.pnl or 0.0) for t in losing_trades))
    profit_factor = (gross_profit / gross_loss) if gross_loss > 0 else (gross_profit if gross_profit > 0 else 0.0)

    # Max Drawdown from Snapshots
    max_dd = 0.0
    peak = 0.0
    for snap in reversed(snapshots):
        eq = float(snap.equity_egp or 0.0)
        if eq > peak:
            peak = eq
        if peak > 0:
            dd = (peak - eq) / peak * 100.0
            if dd > max_dd:
                max_dd = dd

    # -------------------------------------------------------------------------
    # SHEET 1: Executive Summary
    # -------------------------------------------------------------------------
    ws_summary = wb.create_sheet(title="Executive Summary")
    now_str = TimeUtils.now().strftime("%Y-%m-%d %H:%M:%S")

    ws_summary.cell(row=1, column=1, value="HORUS ANALYTICS II — PORTFOLIO TREASURY AUDIT").font = title_font
    ws_summary.cell(row=2, column=1, value=f"Portfolio: {portfolio.name} (#{portfolio.id}) | Generated: {now_str}").font = subtitle_font

    # Profile & Structure Card
    ws_summary.cell(row=4, column=1, value="PORTFOLIO SPECIFICATION").font = section_font
    ws_summary.cell(row=5, column=1, value="Property").font = header_font
    ws_summary.cell(row=5, column=2, value="Value").font = header_font
    apply_headers(ws_summary, 5, 2)

    meta_rows = [
        ("Portfolio ID", f"#{portfolio.id}"),
        ("Portfolio Name", str(portfolio.name)),
        ("Portfolio Classification", str(portfolio.type or "USER")),
        ("Autonomous Engine Managed", "YES (Strategy Managed)" if portfolio.auto_manage else "NO (Manual / Desk)"),
        ("Description", str(portfolio.description or "Primary Treasury Book")),
        ("Report Currency Baseline", "EGP / USD Dual"),
        ("Active Open Holdings", f"{len(open_positions)} Positions"),
        ("Settled Closed Trades", f"{settled_trades_count} Trades"),
    ]

    for idx, (label, val) in enumerate(meta_rows, start=6):
        c1 = ws_summary.cell(row=idx, column=1, value=label)
        c2 = ws_summary.cell(row=idx, column=2, value=val)
        c1.font = bold_font; c1.alignment = align_left; c1.border = thin_border
        c2.font = regular_font; c2.alignment = align_left; c2.border = thin_border

    # Balance & Capital Card
    ws_summary.cell(row=16, column=1, value="CAPITAL & TREASURY RESERVES").font = section_font
    ws_summary.cell(row=17, column=1, value="Balance Category").font = header_font
    ws_summary.cell(row=17, column=2, value="Amount").font = header_font
    apply_headers(ws_summary, 17, 2)

    cap_rows = [
        ("Total Estimated Net Worth (EGP)", f"{net_worth_egp:,.2f} EGP"),
        ("Free Margin Available (EGP)", f"{cash_egp:,.2f} EGP"),
        ("Free Margin Available (USD)", f"${cash_usd:,.2f} USD"),
        ("Invested Market Value (EGP)", f"{total_market_value:,.2f} EGP"),
        ("Total Cost Basis (EGP)", f"{total_cost_basis:,.2f} EGP"),
        ("Cash Drag Allocation %", f"{cash_drag_pct:.1f}%"),
        ("Invested Capital Allocation %", f"{invested_pct:.1f}%"),
    ]

    for idx, (label, val) in enumerate(cap_rows, start=18):
        c1 = ws_summary.cell(row=idx, column=1, value=label)
        c2 = ws_summary.cell(row=idx, column=2, value=val)
        c1.font = bold_font; c1.alignment = align_left; c1.border = thin_border
        c2.font = bold_font if "Net Worth" in label else regular_font
        c2.alignment = align_right; c2.border = thin_border

    # Performance & Risk Card
    ws_summary.cell(row=4, column=4, value="PERFORMANCE & RISK TELEMETRY").font = section_font
    ws_summary.cell(row=5, column=4, value="Telemetry Metric").font = header_font
    ws_summary.cell(row=5, column=5, value="Score / Value").font = header_font
    apply_headers(ws_summary, 5, 2, fill=sub_header_fill)
    for col in (4, 5):
        ws_summary.cell(row=5, column=col).fill = sub_header_fill

    perf_rows = [
        ("Floating MTM PnL (Open Positions)", f"{floating_mtm_pnl:+,.2f} EGP", green_font if floating_mtm_pnl >= 0 else red_font),
        ("Settled Realized PnL (Closed Trades)", f"{settled_realized_pnl:+,.2f} EGP", green_font if settled_realized_pnl >= 0 else red_font),
        ("Win Rate % (Closed Trades)", f"{win_rate:.1f}%", green_font if win_rate >= 50 else amber_font),
        ("Profit Factor", f"{profit_factor:.2f}", green_font if profit_factor >= 1.5 else regular_font),
        ("Winning Trades Count", str(len(winning_trades)), regular_font),
        ("Losing Trades Count", str(len(losing_trades)), regular_font),
        ("Max Historical Drawdown", f"{max_dd:.2f}%", red_font if max_dd > 15 else regular_font),
    ]

    for idx, (label, val, f_style) in enumerate(perf_rows, start=6):
        c1 = ws_summary.cell(row=idx, column=4, value=label)
        c2 = ws_summary.cell(row=idx, column=5, value=val)
        c1.font = bold_font; c1.alignment = align_left; c1.border = thin_border
        c2.font = f_style; c2.alignment = align_right; c2.border = thin_border

    auto_fit_columns(ws_summary, max_cols=6)

    # -------------------------------------------------------------------------
    # SHEET 2: Holdings & Valuation
    # -------------------------------------------------------------------------
    ws_holdings = wb.create_sheet(title="Holdings & Valuation")
    holding_headers = [
        "Ticker", "Currency", "Sector", "Shares", "Avg Entry Price",
        "Live Market Price", "Cost Basis", "Current Market Value",
        "Unrealized MTM PnL", "Return %", "Stop Loss", "Distance to SL %",
        "Take Profit 1", "Take Profit 2", "TP1 Hit", "Risk Status", "Notes"
    ]
    ws_holdings.append(holding_headers)
    apply_headers(ws_holdings, 1, len(holding_headers))

    for idx, pos in enumerate(open_positions, start=2):
        ticker = str(pos.ticker or "").upper()
        curr = str(pos.currency or "EGP").upper()
        sec = str(pos.sector or "-")
        shares = int(pos.shares or 0)
        entry = float(pos.entry_price or 0.0)
        live = float(pos.current_price if pos.current_price is not None else entry)
        cost = entry * shares
        mval = live * shares
        unrealized = mval - cost
        ret_pct = ((live - entry) / entry) if entry > 0 else 0.0
        sl = float(pos.stop_loss) if pos.stop_loss is not None else None
        tp1 = float(pos.target_price) if pos.target_price is not None else None
        tp2 = float(pos.target_price_2) if pos.target_price_2 is not None else None
        tp1_hit_label = "YES" if pos.tp1_hit else "NO"
        risk_status = str(pos.status or "OPEN").upper()
        notes = str(pos.notes or "-")

        dist_sl_pct = ((live - sl) / live) if sl and live > 0 else None

        row_vals = [
            ticker, curr, sec, shares, entry,
            live, cost, mval,
            unrealized, ret_pct, sl if sl else "-", dist_sl_pct if dist_sl_pct is not None else "-",
            tp1 if tp1 else "-", tp2 if tp2 else "-", tp1_hit_label, risk_status, notes
        ]
        ws_holdings.append(row_vals)

        for col_idx in range(1, len(holding_headers) + 1):
            cell = ws_holdings.cell(row=idx, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = align_center

            # Format numbers
            if col_idx in (5, 6, 7, 8, 9, 11, 13, 14):  # Prices, Values, PnL
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"
                    cell.alignment = align_right
            elif col_idx == 4:  # Shares
                cell.number_format = "#,##0"
                cell.alignment = align_right
            elif col_idx in (10, 12):  # Return %, Distance to SL %
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "+0.00%;-0.00%;0.00%"
                    cell.alignment = align_right

        # Conditional formatting for Unrealized & Return %
        pnl_cell = ws_holdings.cell(row=idx, column=9)
        ret_cell = ws_holdings.cell(row=idx, column=10)
        if unrealized >= 0:
            pnl_cell.fill = green_fill; pnl_cell.font = green_font
            ret_cell.fill = green_fill; ret_cell.font = green_font
        else:
            pnl_cell.fill = red_fill; pnl_cell.font = red_font
            ret_cell.fill = red_fill; ret_cell.font = red_font

    auto_fit_columns(ws_holdings)

    # -------------------------------------------------------------------------
    # SHEET 3: Trade Journal
    # -------------------------------------------------------------------------
    ws_trades = wb.create_sheet(title="Trade Journal")
    trade_headers = [
        "Trade ID", "Ticker", "Currency", "Shares", "Entry Date",
        "Entry Price", "Exit Date", "Exit Price", "Realized PnL",
        "Return %", "Close Reason", "Lane Attribution"
    ]
    ws_trades.append(trade_headers)
    apply_headers(ws_trades, 1, len(trade_headers))

    for idx, trade in enumerate(trades, start=2):
        t_id = trade.id
        ticker = str(trade.ticker or "").upper()
        curr = str(trade.currency or "EGP").upper()
        shares = int(trade.shares or 0)
        entry_d = trade.entry_date.strftime("%Y-%m-%d %H:%M") if trade.entry_date else "-"
        entry_p = float(trade.entry_price or 0.0)
        exit_d = trade.exit_date.strftime("%Y-%m-%d %H:%M") if trade.exit_date else "-"
        exit_p = float(trade.exit_price or 0.0)
        pnl = float(trade.pnl or 0.0)
        pnl_pct = (float(trade.pnl_pct or 0.0) / 100.0) if trade.pnl_pct is not None else (((exit_p - entry_p) / entry_p) if entry_p > 0 else 0.0)
        reason = str(trade.reason or "MANUAL").upper()

        exec_row = execution_by_trade_id.get(t_id)
        lane = str(exec_row.trigger_source if exec_row else "MANUAL").upper()

        row_vals = [
            t_id, ticker, curr, shares, entry_d,
            entry_p, exit_d, exit_p, pnl,
            pnl_pct, reason, lane
        ]
        ws_trades.append(row_vals)

        for col_idx in range(1, len(trade_headers) + 1):
            cell = ws_trades.cell(row=idx, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = align_center

            if col_idx in (6, 8, 9):
                cell.number_format = "#,##0.00"
                cell.alignment = align_right
            elif col_idx == 4:
                cell.number_format = "#,##0"
                cell.alignment = align_right
            elif col_idx == 10:
                cell.number_format = "+0.00%;-0.00%;0.00%"
                cell.alignment = align_right

        pnl_cell = ws_trades.cell(row=idx, column=9)
        ret_cell = ws_trades.cell(row=idx, column=10)
        if pnl >= 0:
            pnl_cell.fill = green_fill; pnl_cell.font = green_font
            ret_cell.fill = green_fill; ret_cell.font = green_font
        else:
            pnl_cell.fill = red_fill; pnl_cell.font = red_font
            ret_cell.fill = red_fill; ret_cell.font = red_font

    auto_fit_columns(ws_trades)

    # -------------------------------------------------------------------------
    # SHEET 4: Equity History
    # -------------------------------------------------------------------------
    ws_snaps = wb.create_sheet(title="Equity History")
    snap_headers = [
        "Snapshot Date", "Equity (EGP)", "Equity (USD)", "Cash Reserve (EGP)",
        "Cash Reserve (USD)", "Open Positions Count"
    ]
    ws_snaps.append(snap_headers)
    apply_headers(ws_snaps, 1, len(snap_headers))

    for idx, snap in enumerate(snapshots, start=2):
        s_date = snap.date.strftime("%Y-%m-%d") if snap.date else "-"
        eq_egp = float(snap.equity_egp or 0.0)
        eq_usd = float(snap.equity_usd or 0.0)
        c_egp = float(snap.cash_egp or 0.0)
        c_usd = float(snap.cash_usd or 0.0)
        p_count = int(snap.position_count or 0)

        row_vals = [s_date, eq_egp, eq_usd, c_egp, c_usd, p_count]
        ws_snaps.append(row_vals)

        for col_idx in range(1, len(snap_headers) + 1):
            cell = ws_snaps.cell(row=idx, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = align_center

            if col_idx in (2, 3, 4, 5):
                cell.number_format = "#,##0.00"
                cell.alignment = align_right
            elif col_idx == 6:
                cell.number_format = "#,##0"
                cell.alignment = align_right

    auto_fit_columns(ws_snaps)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
