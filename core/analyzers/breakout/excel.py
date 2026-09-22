from __future__ import annotations

import os
from typing import Any
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core import TimeUtils
from core.settings import settings
from .pivots import MAX_SPREAD_PCT


def export_scan_results_to_excel(
    df_results: pd.DataFrame,
    timestamp: str | None = None,
    date_str: str | None = None,
) -> str | None:
    """Save formatted Excel report and print categorized summary to console."""
    if df_results is None or df_results.empty:
        print("❌ No valid stocks found to analyze.")
        return None

    if timestamp is None:
        timestamp = TimeUtils.now().strftime("%H%M")
    if date_str is None:
        date_str = TimeUtils.now().strftime("%d-%m-%y")

    reports_folder = os.path.join(settings.REPORTS_DIR, f"Reports_{date_str}")
    if not os.path.exists(reports_folder):
        os.makedirs(reports_folder, exist_ok=True)

    filename = os.path.join(reports_folder, f"Full_Market_Scan_{timestamp}.xlsx")

    try:
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "🔥 Loki Market Scan"

        headers = list(df_results.columns)
        ws.append(headers)

        header_fill = PatternFill(start_color="1A1A2E", end_color="1A1A2E", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)

        score_colors = {
            10: "006400", 9: "228B22", 8: "32CD32",
            7: "7FFF00", 6: "9ACD32", 5: "FFD700",
            4: "FFA500", 3: "FF8C00", 2: "FF6347",
            1: "FF0000", 0: "8B0000"
        }

        loki_yes_fill = PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid")
        loki_yes_font = Font(bold=True, color="000000", size=11)
        loki_no_fill = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")

        bullish_font = Font(color="006400", bold=True, size=10)
        bearish_font = Font(color="8B0000", bold=True, size=10)

        status_styles = {
            "HIGH CONVICTION BUY": (PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid"),
                                    Font(color="000000", bold=True, size=11)),
            "STRONG BUY": (PatternFill(start_color="32CD32", end_color="32CD32", fill_type="solid"),
                           Font(color="FFFFFF", bold=True, size=10)),
            "MODERATE BUY": (PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid"),
                             Font(color="FFFFFF", bold=True)),
            "INSTITUTIONAL ACTIVITY": (PatternFill(start_color="9370DB", end_color="9370DB", fill_type="solid"),
                                       Font(color="FFFFFF", bold=True)),
            "BREAKOUT - CONFIRMED": (PatternFill(start_color="4169E1", end_color="4169E1", fill_type="solid"),
                                     Font(color="FFFFFF", bold=True)),
            "BREAKOUT": (PatternFill(start_color="87CEEB", end_color="87CEEB", fill_type="solid"),
                         Font(color="000000", bold=True)),
            "WATCHLIST - NEAR RESISTANCE": (PatternFill(start_color="F0E68C", end_color="F0E68C", fill_type="solid"),
                                           Font(color="000000", bold=True)),
            "POTENTIAL MANIPULATION": (PatternFill(start_color="F43F5E", end_color="F43F5E", fill_type="solid"),
                                       Font(color="FFFFFF", bold=True)),
            "LIQUIDITY TRAP": (PatternFill(start_color="E11D48", end_color="E11D48", fill_type="solid"),
                               Font(color="FFFFFF", bold=True)),
            "ACCUMULATION PHASE": (PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid"),
                                  Font(color="666666"))
        }

        cp_fill = PatternFill(start_color="FFF8DC", end_color="FFF8DC", fill_type="solid")
        cp1_fill = PatternFill(start_color="FFE4B5", end_color="FFE4B5", fill_type="solid")
        cp2_fill = PatternFill(start_color="FFE4E1", end_color="FFE4E1", fill_type="solid")
        tp_fill = PatternFill(start_color="E0FFE0", end_color="E0FFE0", fill_type="solid")

        thin_border = Border(
            left=Side(style='thin', color='B0B0B0'),
            right=Side(style='thin', color='B0B0B0'),
            top=Side(style='thin', color='B0B0B0'),
            bottom=Side(style='thin', color='B0B0B0')
        )

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(
                left=Side(style='medium', color='000000'),
                right=Side(style='medium', color='000000'),
                top=Side(style='medium', color='000000'),
                bottom=Side(style='medium', color='000000')
            )

        col_map = {col: i for i, col in enumerate(headers)}
        result_records = df_results.to_dict(orient='records')
        for idx, row in enumerate(result_records, start=2):
            ws.append([row.get(header) for header in headers])

            for cell in ws[idx]:
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")

            if 'Signal_Score' in col_map:
                score_cell = ws[idx][col_map['Signal_Score']]
                score = int(row['Signal_Score'])
                if score in score_colors:
                    score_cell.fill = PatternFill(start_color=score_colors[score],
                                                  end_color=score_colors[score], fill_type="solid")
                    score_cell.font = Font(bold=True, color="FFFFFF" if score >= 5 else "000000", size=12)

            if 'Institutional_Signal' in col_map:
                loki_cell = ws[idx][col_map['Institutional_Signal']]
                if row['Institutional_Signal'] == "✅":
                    loki_cell.fill = loki_yes_fill
                    loki_cell.font = loki_yes_font
                else:
                    loki_cell.fill = loki_no_fill

            if 'Status' in col_map:
                status_cell = ws[idx][col_map['Status']]
                status_val = str(status_cell.value)
                for key, (fill, font) in status_styles.items():
                    if key in status_val:
                        status_cell.fill = fill
                        status_cell.font = font
                        break
                status_cell.alignment = Alignment(horizontal="left", vertical="center")

            if 'Trend' in col_map:
                trend_cell = ws[idx][col_map['Trend']]
                if trend_cell.value == "BULLISH":
                    trend_cell.font = bullish_font
                elif trend_cell.value == "BEARISH":
                    trend_cell.font = bearish_font

            if 'Resistance_20D' in col_map:
                ws[idx][col_map['Resistance_20D']].fill = cp_fill
                ws[idx][col_map['Resistance_20D']].font = Font(bold=True, color="8B4513")

            if 'Key_Resistance_1' in col_map:
                ws[idx][col_map['Key_Resistance_1']].fill = cp1_fill
                ws[idx][col_map['Key_Resistance_1']].font = Font(bold=True, color="D2691E")

            if 'Key_Resistance_2' in col_map:
                ws[idx][col_map['Key_Resistance_2']].fill = cp2_fill
                ws[idx][col_map['Key_Resistance_2']].font = Font(bold=True, color="CD853F")

            for dist_col in ['Distance_%', 'Key_Res_1_Dist_%', 'Key_Res_2_Dist_%']:
                if dist_col in col_map:
                    dist_cell = ws[idx][col_map[dist_col]]
                    dist_cell.number_format = '0.00"%"'
                    try:
                        dist_val = float(row[dist_col])
                        if dist_val > 0:
                            dist_cell.font = Font(color="006400", bold=True)
                        elif dist_val < 0:
                            dist_cell.font = Font(color="8B0000", bold=True)
                    except:
                        pass

            for tp_col in ['Stop_Loss', 'Target_1', 'Target_2', 'Target_3', 'Target_4']:
                if tp_col in col_map:
                    tp_cell = ws[idx][col_map[tp_col]]
                    tp_cell.fill = tp_fill
                    if 'SL' in tp_col:
                        tp_cell.font = Font(color="8B0000", bold=True)
                    else:
                        tp_cell.font = Font(color="006400", bold=True)
                    tp_cell.number_format = '0.00'

            if 'RSI' in col_map:
                rsi_cell = ws[idx][col_map['RSI']]
                rsi_cell.number_format = '0.0'
                try:
                    rsi_val = float(row['RSI'])
                    if rsi_val >= 70:
                        rsi_cell.font = Font(color="8B0000", bold=True)
                    elif rsi_val <= 30:
                        rsi_cell.font = Font(color="006400", bold=True)
                    elif 50 <= rsi_val < 70:
                        rsi_cell.font = Font(color="228B22")
                except:
                    pass

            for price_col in ['Price', 'EMA9', 'ATR']:
                if price_col in col_map:
                    ws[idx][col_map[price_col]].number_format = '0.00'

            if 'Rel_Volume' in col_map:
                ws[idx][col_map['Rel_Volume']].number_format = '0.00"x"'

            if 'Price_Move_%' in col_map:
                ws[idx][col_map['Price_Move_%']].number_format = '0.00"%"'

            if 'Avg_Turnover_M' in col_map:
                ws[idx][col_map['Avg_Turnover_M']].number_format = '0.00"M"'

            if 'Estimated_Spread_%' in col_map:
                spread_cell = ws[idx][col_map['Estimated_Spread_%']]
                spread_cell.number_format = '0.00"%"'
                try:
                    spread_val = float(row['Estimated_Spread_%'])
                    if spread_val > MAX_SPREAD_PCT * 100:
                        spread_cell.font = Font(color="8B0000", bold=True)
                    else:
                        spread_cell.font = Font(color="006400")
                except:
                    pass

            if 'Risk_Reward_Ratio' in col_map:
                rr_cell = ws[idx][col_map['Risk_Reward_Ratio']]
                rr_cell.number_format = '0.00":1"'
                try:
                    rr_val = float(row['Risk_Reward_Ratio'])
                    if rr_val >= 2:
                        rr_cell.font = Font(color="006400", bold=True)
                except:
                    pass

        column_widths = {
            'Ticker': 10,
            'Currency': 10,
            'Status': 35,
            'Signal_Score': 8,
            'Signal_Reasons': 25,
            'Trend': 10,
            'Price': 10,
            'EMA9': 10,
            'RSI': 8,
            'Rel_Volume': 10,
            'Price_Move_%': 11,
            'Avg_Turnover_M': 12,
            'Estimated_Spread_%': 18,
            'Institutional_Signal': 10,
            'Resistance_20D': 13,
            'Key_Resistance_1': 13,
            'Key_Resistance_2': 13,
            'Distance_%': 11,
            'Key_Res_1_Dist_%': 13,
            'Key_Res_2_Dist_%': 13,
            'ATR': 8,
            'Force_Power': 11,
            'Stop_Loss': 10,
            'Target_1': 10,
            'Target_2': 10,
            'Target_3': 12,
            'Target_4': 14,
            'Risk_Reward_Ratio': 10
        }

        for col_name, width in column_widths.items():
            if col_name in col_map:
                col_letter = get_column_letter(col_map[col_name] + 1)
                ws.column_dimensions[col_letter].width = width

        ws.freeze_panes = 'B2'
        ws.row_dimensions[1].height = 25
        wb.save(filename)

        print("=" * 100)
        print(f"{'💎 MARKET ANALYSIS REPORT GENERATED':^100}")
        print("=" * 100)
        print(f"📄 File: {filename}")
        print(f"📊 Total Stocks Analyzed: {len(df_results)}")
        print("=" * 100)

        status_categories = [
            ("🐋 HIGH CONVICTION BUY", "HIGH CONVICTION BUYS"),
            ("🔥 STRONG BUY", "🔥 STRONG BUYS"),
            ("⚡ MODERATE BUY", "⚡ MODERATE BUYS"),
            ("🐺 INSTITUTIONAL ACTIVITY", "🐺 INSTITUTIONAL ACTIVITY"),
            ("🔥 BREAKOUT - CONFIRMED", "🔥 BREAKOUT - CONFIRMEDS"),
            ("⚡ BREAKOUT", "WEAK BREAKOUTS"),
            ("⚠️ POTENTIAL MANIPULATION", "⚠️ BLOCKED MANIPULATION PUMPS"),
            ("⚠️ LIQUIDITY TRAP", "⚠️ BLOCKED LIQUIDITY/RETAIL TRAPS"),
            ("👀 WATCHLIST - NEAR RESISTANCE", "👀 WATCHLIST - NEAR RESISTANCE (Near Critical Points)")
        ]

        for status_key, category_title in status_categories:
            category_stocks = df_results[df_results['Status'].str.contains(status_key, regex=False)]
            if len(category_stocks) > 0:
                print(f"\n{'─' * 100}")
                print(f"  {category_title} ({len(category_stocks)} stocks)")
                print(f"{'─' * 100}")

                display_limit = 20 if status_key == "👀 WATCHLIST - NEAR RESISTANCE" else 15
                for row in category_stocks.head(display_limit).to_dict(orient='records'):
                    loki_badge = "🐺" if row['Institutional_Signal'] == "✅" else "  "
                    currency_symbol = "$" if row['Currency'] == "USD" else "E£"

                    ticker_info = f"{row['Ticker']:8} ({row['Currency']:3})"
                    score_info = f"Score: {row['Signal_Score']:2}"
                    price_info = f"{currency_symbol}{row['Price']:>7.2f}"
                    tp_info = f"TP2: {currency_symbol}{row['Target_2']:>7.2f}"
                    rr_info = f"R/R: {row['Risk_Reward_Ratio']:>4.1f}x"
                    rsi_info = f"RSI: {row['RSI']:>5.1f}"

                    print(f"  {loki_badge} {ticker_info} | {score_info} | {price_info} → {tp_info} | {rr_info} | {rsi_info}")

                if len(category_stocks) > display_limit:
                    remaining = len(category_stocks) - display_limit
                    print(f"  {'':2} ... and {remaining} more stocks in this category")

        print(f"\n{'=' * 100}")
        print(f"{'✅ ANALYSIS COMPLETE - Check Excel file for full details':^100}")
        print(f"{'=' * 100}\n")
        return filename

    except Exception as e:
        print(f"❌ Error saving Excel file: {e}")
        csv_filename = filename.replace('.xlsx', '.csv')
        df_results.to_csv(csv_filename, index=False)
        print(f"💾 Saved as CSV instead: {csv_filename}")
        return csv_filename
