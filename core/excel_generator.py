"""
EXCEL REPORT GENERATOR MODULE
================================
Builds structured, styled Excel (.xlsx) workbooks from Market Replay and simulation session states.
Utilizes openpyxl for multi-tab layouts, conditional fills, currency formatting, and auto-adjusted column widths.
"""

import io
import datetime
from typing import Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def build_replay_excel_report(replay_state: dict[str, Any]) -> bytes:
    """Generate a 4-sheet styled Excel workbook from a replay state dictionary.

    Returns:
        bytes: Raw XLSX binary data ready for HTTP response or file output.
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    if wb.active is not None:
        wb.remove(wb.active)

    # Reusable Styling Definitions
    font_family = "Segoe UI"
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    
    title_font = Font(name=font_family, size=16, bold=True, color="0F172A")
    subtitle_font = Font(name=font_family, size=10, italic=True, color="64748B")
    section_font = Font(name=font_family, size=12, bold=True, color="1E293B")
    bold_font = Font(name=font_family, size=10, bold=True, color="0F172A")
    regular_font = Font(name=font_family, size=10, color="334155")
    
    # Outcomes Fills
    green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    green_font = Font(name=font_family, size=10, bold=True, color="15803D")
    
    red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    red_font = Font(name=font_family, size=10, bold=True, color="B91C1C")
    
    blue_fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
    blue_font = Font(name=font_family, size=10, bold=True, color="0369A1")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    thin_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    def apply_header_styles(ws, row_idx: int, col_count: int):
        for col in range(1, col_count + 1):
            cell = ws.cell(row=row_idx, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = thin_border

    def auto_fit_columns(ws, max_cols: int | None = None):
        ws.views.sheetView[0].showGridLines = True
        cols = range(1, (max_cols or ws.max_column) + 1)
        for col in cols:
            max_len = 0
            col_letter = get_column_letter(col)
            for row in range(1, ws.max_row + 1):
                val = ws.cell(row=row, column=col).value
                if val is not None:
                    max_len = max(max_len, len(str(val)))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # -------------------------------------------------------------------------
    # SHEET 1: Executive Summary
    # -------------------------------------------------------------------------
    ws_summary = wb.create_sheet(title="Executive Summary")
    
    ws_summary.cell(row=1, column=1, value="HORUS ANALYTICS II — MARKET REPLAY REPORT").font = title_font
    generated_at_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws_summary.cell(row=2, column=1, value=f"Generated: {generated_at_str} | Session Replay Engine").font = subtitle_font
    
    # Parameters Card
    ws_summary.cell(row=4, column=1, value="SESSION PARAMETERS").font = section_font
    ws_summary.cell(row=5, column=1, value="Parameter").font = header_font
    ws_summary.cell(row=5, column=2, value="Value").font = header_font
    apply_header_styles(ws_summary, 5, 2)

    params = [
        ("Target Replay Date", str(replay_state.get("date") or replay_state.get("start_date") or "N/A")),
        ("Replay Mode", str(replay_state.get("mode") or "SINGLE_DAY")),
        ("Time Speed Multiplier", f"{replay_state.get('speed', 1)}x"),
        ("Market Universe", str(replay_state.get("market") or "EGX30")),
        ("Scanner Profile", str(replay_state.get("profile_name") or "Horus Core")),
        ("Execution Status", str(replay_state.get("status") or "IDLE")),
        ("Simulated Market Hours", f"{replay_state.get('market_open', '10:00')} → {replay_state.get('market_close', '14:30')}"),
        ("Total Completed Ticks", f"{replay_state.get('ticks_completed', 0)} / {replay_state.get('total_ticks', 0)}"),
    ]
    
    for r_offset, (p_name, p_val) in enumerate(params, start=6):
        c1 = ws_summary.cell(row=r_offset, column=1, value=p_name)
        c2 = ws_summary.cell(row=r_offset, column=2, value=p_val)
        c1.font = bold_font; c1.alignment = align_left; c1.border = thin_border
        c2.font = regular_font; c2.alignment = align_left; c2.border = thin_border

    # KPI Metrics Card
    active_trades = list(replay_state.get("active_trades") or [])
    closed_trades = [t for t in active_trades if t.get("state") == "CLOSED"]
    open_trades = [t for t in active_trades if t.get("state") != "CLOSED"]
    tp_hits = [t for t in closed_trades if "target" in str(t.get("exit_reason") or "").lower() or t.get("tp1_hit")]
    stop_hits = [t for t in closed_trades if "stop" in str(t.get("exit_reason") or "").lower()]
    
    win_rate = (len(tp_hits) / len(closed_trades) * 100.0) if closed_trades else 0.0
    
    total_pnl_egp = 0.0
    for t in closed_trades:
        entry = float(t.get("entry_price") or 0)
        exit_p = float(t.get("exit_price") or entry)
        shares = int(t.get("shares") or 0)
        total_pnl_egp += (exit_p - entry) * shares

    unrealized_pnl_egp = 0.0
    for t in open_trades:
        entry = float(t.get("entry_price") or 0)
        current = float(t.get("current_price") or t.get("last_price") or entry)
        shares = int(t.get("shares") or 0)
        unrealized_pnl_egp += (current - entry) * shares

    ws_summary.cell(row=15, column=1, value="KEY PERFORMANCE INDICATORS").font = section_font
    ws_summary.cell(row=16, column=1, value="Metric").font = header_font
    ws_summary.cell(row=16, column=2, value="Value").font = header_font
    apply_header_styles(ws_summary, 16, 2)

    kpis = [
        ("Unique Signals Scanned", len(replay_state.get("unique_signals") or []) or replay_state.get("signals_found", 0)),
        ("Total Trades Opened", len(active_trades)),
        ("Closed Trades", len(closed_trades)),
        ("Open Floating Positions", len(open_trades)),
        ("Target Wins (TP1 / TP2)", len(tp_hits)),
        ("Stop Loss Exits", len(stop_hits)),
        ("Win Rate % (Closed)", f"{win_rate:.1f}%"),
        ("Net Realized PnL (EGP)", f"{total_pnl_egp:,.2f} EGP"),
        ("Floating Unrealized PnL (EGP)", f"{unrealized_pnl_egp:,.2f} EGP"),
        ("Pending Entries Waiting Open", len(replay_state.get("pending_entries") or [])),
    ]

    for r_offset, (k_name, k_val) in enumerate(kpis, start=17):
        c1 = ws_summary.cell(row=r_offset, column=1, value=k_name)
        c2 = ws_summary.cell(row=r_offset, column=2, value=k_val)
        c1.font = bold_font; c1.alignment = align_left; c1.border = thin_border
        c2.font = regular_font; c2.alignment = align_right; c2.border = thin_border
        if k_name == "Net Realized PnL (EGP)":
            c2.font = green_font if total_pnl_egp >= 0 else red_font
        elif k_name == "Floating Unrealized PnL (EGP)":
            c2.font = green_font if unrealized_pnl_egp >= 0 else red_font

    auto_fit_columns(ws_summary, max_cols=3)

    # -------------------------------------------------------------------------
    # SHEET 2: Executed Trades
    # -------------------------------------------------------------------------
    ws_trades = wb.create_sheet(title="Executed Trades")
    
    trade_headers = [
        "Ticker", "Side", "State", "Entry Time", "Entry Price (EGP)", 
        "Shares", "Exit Time", "Exit Price (EGP)", "Exit Reason", "PnL (EGP)", "Return %"
    ]
    ws_trades.append(trade_headers)
    apply_header_styles(ws_trades, 1, len(trade_headers))

    for idx, trade in enumerate(active_trades, start=2):
        ticker = str(trade.get("ticker") or "").upper()
        side = str(trade.get("side") or "BUY").upper()
        state = str(trade.get("state") or "OPEN").upper()
        entry_time = str(trade.get("entry_at") or "-")
        entry_price = float(trade.get("entry_price") or 0.0)
        shares = int(trade.get("shares") or 0)
        exit_time = str(trade.get("exit_at") or "-")
        exit_price = float(trade.get("exit_price") or entry_price) if state == "CLOSED" else entry_price
        exit_reason = str(trade.get("exit_reason") or ("OPEN" if state != "CLOSED" else "FLAT")).upper()
        
        pnl_egp = (exit_price - entry_price) * shares if state == "CLOSED" else 0.0
        pnl_pct = ((exit_price - entry_price) / entry_price) if entry_price > 0 and state == "CLOSED" else 0.0

        row_vals = [
            ticker, side, state, entry_time, entry_price, 
            shares, exit_time, exit_price if state == "CLOSED" else "-", exit_reason, pnl_egp, pnl_pct
        ]
        ws_trades.append(row_vals)

        # Style row cells
        for col_idx in range(1, len(trade_headers) + 1):
            cell = ws_trades.cell(row=idx, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = align_center

            # Format numbers
            if col_idx in (5, 8, 10):  # Prices & PnL EGP
                cell.number_format = "#,##0.00"
                cell.alignment = align_right
            elif col_idx == 6:  # Shares
                cell.number_format = "#,##0"
                cell.alignment = align_right
            elif col_idx == 11:  # Return %
                cell.number_format = "+0.00%;-0.00%;0.00%"
                cell.alignment = align_right

        # Conditional color highlights for state/outcome
        state_cell = ws_trades.cell(row=idx, column=3)
        pnl_cell = ws_trades.cell(row=idx, column=10)
        ret_cell = ws_trades.cell(row=idx, column=11)
        
        if state == "CLOSED":
            if pnl_egp >= 0:
                state_cell.fill = green_fill; state_cell.font = green_font
                pnl_cell.fill = green_fill; pnl_cell.font = green_font
                ret_cell.fill = green_fill; ret_cell.font = green_font
            else:
                state_cell.fill = red_fill; state_cell.font = red_font
                pnl_cell.fill = red_fill; pnl_cell.font = red_font
                ret_cell.fill = red_fill; ret_cell.font = red_font
        else:
            state_cell.fill = blue_fill; state_cell.font = blue_font

    auto_fit_columns(ws_trades)

    # -------------------------------------------------------------------------
    # SHEET 3: Signals & Intake
    # -------------------------------------------------------------------------
    ws_signals = wb.create_sheet(title="Signals & Intake")
    
    signal_headers = [
        "Ticker", "Score", "Signal Type", "Planned Entry", "Stop Loss", 
        "Target 1", "Target 2", "RSI", "Vol Spike (x)", "Rationale", "Status"
    ]
    ws_signals.append(signal_headers)
    apply_header_styles(ws_signals, 1, len(signal_headers))

    # Collect signals across scan results
    scan_results = list(replay_state.get("scan_results") or [])
    all_signals = []
    seen_sig_keys = set()
    
    for sr in scan_results:
        for s in sr.get("signals") or []:
            if isinstance(s, dict) and s.get("ticker"):
                key = (s.get("ticker"), s.get("type"))
                if key not in seen_sig_keys:
                    seen_sig_keys.add(key)
                    all_signals.append(s)

    for idx, s in enumerate(all_signals, start=2):
        ticker = str(s.get("ticker") or "").upper()
        score = float(s.get("score") or 0.0)
        sig_type = str(s.get("type") or "BUY").upper()
        entry_p = float(s.get("price") or s.get("planned_entry_price") or 0.0)
        sl_p = float(s.get("stop_loss") or 0.0)
        tp1_p = float(s.get("tp1") or s.get("target_price") or 0.0)
        tp2_p = float(s.get("tp2") or s.get("target_price_2") or 0.0)
        rsi_val = float(s.get("rsi")) if s.get("rsi") is not None else "-"
        vol_val = float(s.get("volume_x")) if s.get("volume_x") is not None else "-"
        rationale = str(s.get("rationale") or "-")
        status = str(s.get("pending_action") or s.get("skip_reason") or "SCANNED").upper()

        row_vals = [
            ticker, score, sig_type, entry_p, sl_p, 
            tp1_p, tp2_p, rsi_val, vol_val, rationale, status
        ]
        ws_signals.append(row_vals)

        for col_idx in range(1, len(signal_headers) + 1):
            cell = ws_signals.cell(row=idx, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = align_center

            if col_idx in (4, 5, 6, 7):
                cell.number_format = "#,##0.00"
                cell.alignment = align_right
            elif col_idx == 2:
                cell.number_format = "0.0"

    auto_fit_columns(ws_signals)

    # -------------------------------------------------------------------------
    # SHEET 4: Tick Breakdown
    # -------------------------------------------------------------------------
    ws_ticks = wb.create_sheet(title="Tick Breakdown")
    
    tick_headers = [
        "Tick Index", "Simulated Time", "Scan Label", "Market Regime", 
        "Signals Found", "Skipped Entries", "Pending Remaining"
    ]
    ws_ticks.append(tick_headers)
    apply_header_styles(ws_ticks, 1, len(tick_headers))

    for idx, sr in enumerate(scan_results, start=2):
        t_index = sr.get("tick_index", idx - 1)
        sim_time = str(sr.get("current_time") or "-")
        scan_label = str(sr.get("scan_label") or "INTRADAY").upper()
        regime = str(sr.get("regime") or "NEUTRAL").upper()
        sig_count = len(sr.get("signals") or [])
        skipped_count = int(sr.get("skipped_entries") or 0)
        pending_rem = int((sr.get("pending_entries") or {}).get("remaining") or 0)

        row_vals = [t_index, sim_time, scan_label, regime, sig_count, skipped_count, pending_rem]
        ws_ticks.append(row_vals)

        for col_idx in range(1, len(tick_headers) + 1):
            cell = ws_ticks.cell(row=idx, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = align_center

    auto_fit_columns(ws_ticks)

    # Return raw binary XLSX
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
