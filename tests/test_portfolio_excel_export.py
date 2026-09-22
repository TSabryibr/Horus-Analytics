import io
import pytest
import openpyxl
from datetime import datetime

import database
from core.reports.portfolio_excel import build_portfolio_excel_report
from core.portfolio.import_export import build_portfolio_excel_export_response


def test_build_portfolio_excel_report():
    # 1. Create test portfolio
    p = database.Portfolio.create(
        name="Apex Alpha Fund",
        type="USER",
        auto_manage=True,
        cash_egp=250000.0,
        cash_usd=5000.0,
        description="Core quantitative swing trading portfolio",
    )

    # 2. Add open positions
    database.Position.create(
        portfolio=p.id,
        ticker="COMI",
        shares=1000,
        entry_price=100.0,
        current_price=108.5,
        stop_loss=95.0,
        target_price=115.0,
        target_price_2=120.0,
        tp1_hit=False,
        currency="EGP",
        sector="Banking",
        status="OPEN",
        notes="High momentum breakout",
    )
    database.Position.create(
        portfolio=p.id,
        ticker="ABUK",
        shares=500,
        entry_price=80.0,
        current_price=76.0,
        stop_loss=74.0,
        target_price=90.0,
        currency="EGP",
        sector="Basic Resources",
        status="OPEN",
    )

    # 3. Add closed trades
    database.Trade.create(
        portfolio=p.id,
        ticker="ETEL",
        shares=400,
        entry_price=35.0,
        exit_price=39.2,
        entry_date=datetime(2026, 1, 10, 10, 0),
        exit_date=datetime(2026, 1, 15, 14, 30),
        pnl=1680.0,
        pnl_pct=12.0,
        reason="TP1_HIT",
        currency="EGP",
    )

    # 4. Add historical snapshot
    database.PortfolioSnapshot.create(
        portfolio=p.id,
        date=datetime(2026, 1, 14).date(),
        equity_egp=390000.0,
        equity_usd=5000.0,
        cash_egp=250000.0,
        cash_usd=5000.0,
        position_count=2,
    )

    # Generate workbook bytes
    raw_xlsx = build_portfolio_excel_report(p.id)
    assert raw_xlsx is not None
    assert len(raw_xlsx) > 0

    # Load in openpyxl and verify sheets
    wb = openpyxl.load_workbook(io.BytesIO(raw_xlsx), data_only=True)
    assert set(wb.sheetnames) == {
        "Executive Summary",
        "Holdings & Valuation",
        "Trade Journal",
        "Equity History",
    }

    # Verify Executive Summary content
    ws_sum = wb["Executive Summary"]
    assert "HORUS ANALYTICS II — PORTFOLIO TREASURY AUDIT" in str(ws_sum.cell(row=1, column=1).value)
    
    # Check Holdings Sheet
    ws_hold = wb["Holdings & Valuation"]
    assert ws_hold.cell(row=1, column=1).value == "Ticker"
    assert ws_hold.cell(row=2, column=1).value in ("COMI", "ABUK")

    # Check Trade Journal Sheet
    ws_trades = wb["Trade Journal"]
    assert ws_trades.cell(row=1, column=2).value == "Ticker"
    assert ws_trades.cell(row=2, column=2).value == "ETEL"

    # Check Equity History Sheet
    ws_snaps = wb["Equity History"]
    assert ws_snaps.cell(row=1, column=1).value == "Snapshot Date"


def test_build_portfolio_excel_export_response():
    p = database.Portfolio.create(
        name="Growth Portfolio",
        type="USER",
        cash_egp=100000.0,
        cash_usd=0.0,
    )

    response = build_portfolio_excel_export_response(
        portfolio_id=p.id,
        resolve_portfolio_id_fn=lambda pid: pid,
        select_first_portfolio_id_fn=lambda: p.id,
    )

    assert response.status_code == 200
    assert response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment; filename=\"portfolio_audit_" in response.headers["Content-Disposition"]
    assert response.headers["Content-Disposition"].endswith('.xlsx"')
