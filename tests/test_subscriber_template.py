import io
import pytest
import openpyxl

from core.portfolio.template_generator import (
    build_subscriber_intake_excel_template,
    build_subscriber_intake_csv_template,
    parse_subscriber_intake_file,
)


def test_build_subscriber_intake_excel_template():
    raw_xlsx = build_subscriber_intake_excel_template()
    assert raw_xlsx is not None
    assert len(raw_xlsx) > 0

    wb = openpyxl.load_workbook(io.BytesIO(raw_xlsx), data_only=True)
    assert set(wb.sheetnames) == {"Instructions", "Holdings Intake"}

    ws_intake = wb["Holdings Intake"]
    headers = [cell.value for cell in ws_intake[1]]
    assert "Ticker*" in headers
    assert "Shares*" in headers
    assert "Entry_Price*" in headers


def test_build_subscriber_intake_csv_template():
    csv_text = build_subscriber_intake_csv_template()
    assert "Ticker,Shares,Entry_Price" in csv_text
    assert "COMI" in csv_text


def test_parse_subscriber_intake_excel():
    # 1. Create a simulated user-filled Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Holdings Intake"
    ws.append(["Ticker", "Shares", "Entry_Price", "Stop_Loss", "Target_Price_1", "Target_Price_2", "Currency", "Sector", "Notes"])
    ws.append(["SWDY", 1500, 45.2, 42.0, 50.0, 53.0, "EGP", "Industrial Goods", "Client position 1"])
    ws.append(["ISPH", 3000, 3.8, 3.5, 4.2, 4.6, "EGP", "Healthcare", "Client position 2"])

    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()

    parsed = parse_subscriber_intake_file(file_bytes, "client_holdings.xlsx")
    assert len(parsed) == 2
    assert parsed[0]["ticker"] == "SWDY"
    assert parsed[0]["shares"] == 1500.0
    assert parsed[0]["entry_price"] == 45.2
    assert parsed[0]["stop_loss"] == 42.0
    assert parsed[0]["target_price"] == 50.0
    assert parsed[0]["target_price_2"] == 53.0
    assert parsed[0]["currency"] == "EGP"
    assert parsed[0]["sector"] == "Industrial Goods"
    assert parsed[0]["notes"] == "Client position 1"

    assert parsed[1]["ticker"] == "ISPH"
    assert parsed[1]["shares"] == 3000.0


def test_parse_subscriber_intake_csv():
    csv_data = (
        "Ticker,Shares,Entry_Price,Stop_Loss,Target_Price_1,Currency,Sector,Notes\n"
        "EKHO,800,1.25,1.15,1.45,USD,Energy,USD Holding\n"
        "TMGH,1200,62.0,58.0,70.0,EGP,Real Estate,Core\n"
    ).encode("utf-8")

    parsed = parse_subscriber_intake_file(csv_data, "client_holdings.csv")
    assert len(parsed) == 2
    assert parsed[0]["ticker"] == "EKHO"
    assert parsed[0]["shares"] == 800.0
    assert parsed[0]["currency"] == "USD"
    assert parsed[1]["ticker"] == "TMGH"
    assert parsed[1]["entry_price"] == 62.0


def test_parse_subscriber_intake_with_ticker_aliases_and_currency_strings():
    csv_data = (
        "Ticker,Shares,Entry_Price,Stop_Loss,Target_Price_1,Currency\n"
        "CIB, 500 shares, 104.50 EGP, 98.00 LE, 115.00 EGP, EGP\n"
        'TMG, "1,000", 62.50, 58.00, 72.00, EGP\n'
        "EFG HERMES, 2500, 24.50, 22.00, 28.00, EGP\n"
    ).encode("utf-8")

    parsed = parse_subscriber_intake_file(csv_data, "dirty_client_intake.csv")
    assert len(parsed) == 3
    # Ticker aliases
    assert parsed[0]["ticker"] == "COMI"
    assert parsed[0]["shares"] == 500.0
    assert parsed[0]["entry_price"] == 104.5
    assert parsed[0]["stop_loss"] == 98.0
    assert parsed[0]["target_price"] == 115.0

    assert parsed[1]["ticker"] == "TMGH"
    assert parsed[1]["shares"] == 1000.0

    assert parsed[2]["ticker"] == "EFGD"
    assert parsed[2]["shares"] == 2500.0


def test_subscriber_import_command_sandbox_mode():
    from core.portfolio.management import subscriber_import_command

    csv_data = (
        "Ticker,Shares,Entry_Price,Stop_Loss,Target_Price_1\n"
        "COMI,500,104.50,98.00,115.00\n"
        "ABUK,1000,60.00,55.00,68.00\n"
    ).encode("utf-8")

    result = subscriber_import_command(
        file_bytes=csv_data,
        filename="sandbox_test.csv",
        mode="sandbox",
        portfolio_name="Dr. Hany Sandbox",
        starting_cash_egp=50000.0,
    )

    assert result["status"] == "success"
    assert result["mode"] == "sandbox"
    assert result["portfolio_id"] is None
    assert result["portfolio_name"] == "Dr. Hany Sandbox"
    assert "report" in result
    assert result["report"]["summary"]["open_positions"] == 2
    assert result["report"]["risk"]["health_score"] > 0
    assert len(result["holdings"]) == 2


def test_import_portfolio_csv_bytes_supports_subscriber_template_excel():
    from database import Portfolio, Position
    from core.portfolio.import_export import import_portfolio_csv_bytes

    portfolio = Portfolio.create(name="Subscriber Excel Test", type="USER")

    # Build an in-memory Excel file
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Holdings Intake"
    ws.append(["Ticker", "Shares", "Entry_Price", "Stop_Loss", "Target_Price_1", "Currency", "Sector"])
    ws.append(["CIB", 800, 105.0, 99.0, 115.0, "EGP", "Banking"])
    ws.append(["TMG", 1200, 60.0, 55.0, 70.0, "EGP", "Real Estate"])

    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()

    result = import_portfolio_csv_bytes(
        raw=file_bytes,
        portfolio_id=portfolio.id,
        replace_existing=True,
    )

    assert result["status"] == "success"
    assert result["imported"]["positions"] == 2

    positions = list(Position.select().where(Position.portfolio == portfolio.id))
    assert len(positions) == 2
    tickers = {p.ticker for p in positions}
    assert "COMI" in tickers  # Ticker alias CIB -> COMI
    assert "TMGH" in tickers  # Ticker alias TMG -> TMGH


def test_import_portfolio_csv_bytes_supports_subscriber_template_csv():
    from database import Portfolio, Position
    from core.portfolio.import_export import import_portfolio_csv_bytes

    portfolio = Portfolio.create(name="Subscriber CSV Test", type="USER")

    csv_data = (
        "Ticker,Shares,Entry_Price,Stop_Loss,Target_Price_1,Currency\n"
        "SWDY,1500,45.0,42.0,50.0,EGP\n"
        "ABUK,700,80.0,75.0,90.0,EGP\n"
    ).encode("utf-8")

    result = import_portfolio_csv_bytes(
        raw=csv_data,
        portfolio_id=portfolio.id,
        replace_existing=True,
    )

    assert result["status"] == "success"
    assert result["imported"]["positions"] == 2

    positions = list(Position.select().where(Position.portfolio == portfolio.id))
    assert len(positions) == 2
    tickers = {p.ticker for p in positions}
    assert tickers == {"SWDY", "ABUK"}


