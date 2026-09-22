"""
UNIT TESTS — EXCEL GENERATOR MODULE
=====================================
Tests build_replay_excel_report() to ensure valid openpyxl workbooks are created
with correct multi-tab structures, styling, formatting, and data content.
"""

import io
import openpyxl
from core.excel_generator import build_replay_excel_report


def test_build_replay_excel_report_structure():
    sample_state = {
        "status": "COMPLETED",
        "date": "2026-07-19",
        "speed": 50,
        "mode": "SINGLE_DAY",
        "market": "EGX30",
        "profile_name": "Horus Core",
        "market_open": "10:00",
        "market_close": "14:30",
        "ticks_completed": 54,
        "total_ticks": 54,
        "unique_signals": {"COMI", "SWDY"},
        "signals_found": 2,
        "active_trades": [
            {
                "ticker": "COMI",
                "side": "BUY",
                "state": "CLOSED",
                "entry_price": 82.5,
                "shares": 100,
                "exit_price": 86.1,
                "exit_reason": "target_1",
                "entry_at": "2026-07-19T10:05:00",
                "exit_at": "2026-07-19T11:30:00",
                "tp1_hit": True,
            },
            {
                "ticker": "SWDY",
                "side": "BUY",
                "state": "OPEN",
                "entry_price": 45.2,
                "shares": 200,
                "stop_loss": 43.8,
                "tp1": 48.0,
                "tp2": 50.0,
                "entry_at": "2026-07-19T10:15:00",
            },
        ],
        "pending_entries": [],
        "scan_results": [
            {
                "tick_index": 1,
                "current_time": "10:05",
                "scan_label": "INTRADAY",
                "regime": "BULLISH",
                "signals": [
                    {
                        "ticker": "COMI",
                        "score": 8.5,
                        "type": "BUY",
                        "price": 82.5,
                        "stop_loss": 80.0,
                        "tp1": 86.0,
                        "tp2": 89.0,
                        "rsi": 62.4,
                        "volume_x": 2.1,
                        "rationale": "Breakout above resistance",
                        "pending_action": "queued",
                    }
                ],
                "skipped_entries": 0,
                "pending_entries": {"remaining": 0},
            }
        ],
    }

    excel_bytes = build_replay_excel_report(sample_state)
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    # Load with openpyxl to verify valid workbook structure
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames

    assert "Executive Summary" in sheet_names
    assert "Executed Trades" in sheet_names
    assert "Signals & Intake" in sheet_names
    assert "Tick Breakdown" in sheet_names

    # Check Executive Summary content
    ws_sum = wb["Executive Summary"]
    assert ws_sum["A1"].value == "HORUS ANALYTICS II — MARKET REPLAY REPORT"

    # Check Executed Trades content
    ws_trades = wb["Executed Trades"]
    assert ws_trades.cell(row=1, column=1).value == "Ticker"
    assert ws_trades.cell(row=2, column=1).value == "COMI"
    assert ws_trades.cell(row=2, column=3).value == "CLOSED"
    assert ws_trades.cell(row=3, column=1).value == "SWDY"
    assert ws_trades.cell(row=3, column=3).value == "OPEN"

    # Check Signals content
    ws_sigs = wb["Signals & Intake"]
    assert ws_sigs.cell(row=2, column=1).value == "COMI"
    assert ws_sigs.cell(row=2, column=2).value == 8.5


def test_build_replay_excel_report_empty_state():
    empty_state = {
        "status": "IDLE",
        "ticks_completed": 0,
        "total_ticks": 0,
        "active_trades": [],
        "scan_results": [],
    }

    excel_bytes = build_replay_excel_report(empty_state)
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Executive Summary" in wb.sheetnames
